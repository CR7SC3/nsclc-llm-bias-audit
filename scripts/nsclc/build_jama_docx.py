#!/usr/bin/env python3
"""Render manuscript_nsclc_JAMA_NetworkOpen.md to a plain-black Word .docx.

Mirrors build_manuscript_docx.py: pandoc handles structure, tables, embedded
figures, and the numbered reference list; a python-docx post-pass forces every
style and run to black.

Run from repo root:  python scripts/nsclc/build_jama_docx.py
"""
from pathlib import Path
import pypandoc
from docx import Document
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "docs" / "paper1_nsclc" / "manuscript_nsclc_JAMA_NetworkOpen.md"
OUT = ROOT / "docs" / "paper1_nsclc" / "manuscript_nsclc_JAMA_NetworkOpen.docx"
BLACK = RGBColor(0, 0, 0)


def main():
    pypandoc.convert_file(
        str(SRC),
        "docx",
        outputfile=str(OUT),
        extra_args=["--from=markdown", f"--resource-path={ROOT}"],
    )

    doc = Document(str(OUT))
    for style in doc.styles:
        font = getattr(style, "font", None)
        if font is not None:
            try:
                font.color.rgb = BLACK
            except Exception:
                pass
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = BLACK
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = BLACK
    doc.save(str(OUT))
    print(f"Wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
